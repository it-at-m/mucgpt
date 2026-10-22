from io import BytesIO

from bs4 import BeautifulSoup, NavigableString, Tag
from openpyxl import Workbook


def _extract_prose(soup: BeautifulSoup) -> list[str]:
    prose = []
    for element in soup.children:
        if isinstance(element, Tag):
            if element.name == "table":
                continue
            text = element.get_text("\n", strip=True)
        elif isinstance(element, NavigableString):
            text = element.strip()
        else:
            continue
        if text:
            prose.append(text)
    return prose


def _neutralize_formula(value: str) -> str:
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _fill_table(sheet, table: Tag) -> None:
    for row_index, tr in enumerate(table.find_all("tr"), start=1):
        for col_index, cell in enumerate(tr.find_all(["th", "td"]), start=1):
            sheet.cell(
                row=row_index,
                column=col_index,
                value=_neutralize_formula(cell.get_text(strip=True)),
            )


def convert_xlsx(html: str) -> bytes:
    soup = BeautifulSoup(html, "html.parser")
    workbook = Workbook()
    text_sheet = workbook.active
    text_sheet.title = "Text"

    prose = _extract_prose(soup)
    for row_index, paragraph in enumerate(prose, start=1):
        text_sheet.cell(row=row_index, column=1, value=_neutralize_formula(paragraph))
    if not prose:
        workbook.remove(text_sheet)

    for index, table in enumerate(soup.find_all("table"), start=1):
        sheet = workbook.create_sheet(title=f"Tabelle {index}")
        _fill_table(sheet, table)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
